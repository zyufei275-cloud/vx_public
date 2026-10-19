from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path('D:/vx_public/output')
src=root/'杰出学者的成长节奏_公众号初稿.docx'
dst=root/'杰出学者的成长节奏_公众号初稿_重点高亮版.docx'
doc=Document(src)
targets=[
'研究传递出的一个重要启示是：杰出学者存在不同的成长节奏，短期表现需要放进长期轨迹中理解。',
'科研经历至少20年的287位计算机科学家',
'增长最快的类型，未必拥有最长的高产出阶段。',
'多数学者的p-index经过约10—20年积累达到所观察生涯中的峰值',
'同一个人在不同维度上，可能呈现不同节奏。',
'某一阶段论文数量减少，不能直接等同于所有科研表现同步下降。',
'但无法覆盖创新能力、人才培养和成果转化等全部科研贡献。',
'但这项研究尚不足以确定它们各自的因果作用。',
'评价周期应结合学科和任务特点设计',
'为持续积累、表现不断改善的研究者保留后续机会。',
'优秀学者并没有统一的成长时间表。'
]
before=[p.text for p in doc.paragraphs]
found=[]
for p in doc.paragraphs:
    spans=[]
    for t in targets:
        at=p.text.find(t)
        if at>=0:
            spans.append((at,at+len(t))); found.append(t)
    if not spans: continue
    pos=0
    for run in list(p.runs):
        text=run.text
        if not text:
            continue
        start=pos; end=pos+len(text); pos=end
        bounds=sorted({start,end}|{x for a,b in spans for x in (a,b) if start<x<end})
        for a,b in zip(bounds,bounds[1:]):
            new=deepcopy(run._r)
            for child in list(new):
                if child.tag!=qn('w:rPr'): new.remove(child)
            t=OxmlElement('w:t'); t.set(qn('xml:space'),'preserve'); t.text=text[a-start:b-start]; new.append(t)
            if any(a>=x and b<=y for x,y in spans):
                props=new.get_or_add_rPr()
                old=props.find(qn('w:shd'))
                if old is not None: props.remove(old)
                shade=OxmlElement('w:shd'); shade.set(qn('w:val'),'clear'); shade.set(qn('w:fill'),'FFF2CC'); props.append(shade)
                bold=props.find(qn('w:b'))
                if bold is None: bold=OxmlElement('w:b'); props.append(bold)
                bold.set(qn('w:val'),'1')
            run._r.addprevious(new)
        run._r.getparent().remove(run._r)
assert sorted(found)==sorted(targets),(found,targets)
assert before==[p.text for p in doc.paragraphs]
doc.save(dst)
check=Document(dst)
assert before==[p.text for p in check.paragraphs]
assert len(check.inline_shapes)==2
print('Highlighted passages:',len(found),'Images:',len(check.inline_shapes))
