from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('D:/vx_public')
OUT=ROOT/'output'
OUT.mkdir(exist_ok=True)
doc=Document()
s=doc.sections[0]
s.top_margin=Cm(1.8); s.bottom_margin=Cm(1.8)
s.left_margin=Cm(2.1); s.right_margin=Cm(2.1)
s.page_width=Cm(21); s.page_height=Cm(29.7)
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=styles[name]
    st.font.name='Microsoft YaHei'
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    st.font.color.rgb=RGBColor(0,0,0)
normal=styles['Normal']; normal.font.size=Pt(10.5)
normal.paragraph_format.line_spacing=1.25
normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',25),('Heading 1',16),('Heading 2',12)]:
    styles[name].font.size=Pt(size)
    styles[name].font.bold=True
    styles[name].paragraph_format.space_before=Pt(9)
    styles[name].paragraph_format.space_after=Pt(7)
styles['Caption'].font.size=Pt(9)
styles['Caption'].font.bold=False
styles['Caption'].font.italic=False
styles['Subtitle'].font.italic=False
styles['Caption'].paragraph_format.line_spacing=1.2
styles['Caption'].paragraph_format.space_after=Pt(8)

def p(t,bold=False):
    a=doc.add_paragraph(); r=a.add_run(t); r.bold=bold; return a
def h(t): doc.add_heading(t,1)
def sub(t): doc.add_heading(t,2)
def page(): doc.add_page_break()
def pic(n,width=16.5):
    a=doc.add_paragraph(); a.alignment=1
    a.add_run().add_picture(str(ROOT/'analysis_tmp'/'word'/'media'/n),width=Cm(width))
def cap(t): doc.add_paragraph(t,'Caption')

doc.add_paragraph('研究成果推介','Subtitle')
doc.add_paragraph('杰出学者的成长节奏', 'Title')
p('少年得志还是大器晚成？一项对287位计算机科学家的研究，带我们重新理解科研成长。',True)
p('有人很早就崭露头角，有人在漫长积累之后才逐渐受到关注。面对这样的差异，我们常常会问：起步快的人，是否更容易走得远？暂时没有亮眼成绩的人，又该如何看待自己的发展？')
p('对科研工作者而言，这些问题尤其难以用某一年的论文数量回答。一项成果的发表、影响力的形成，以及同行的认可，可能发生在不同的时间。')
p('南京大学信息管理学院Scool实验室的一项研究，将目光放在学者完整的成长轨迹上。研究发表于《Journal of Informetrics》，分析了287位ACM会士的长期文献计量表现。')
p('研究传递出的一个重要启示是：杰出学者存在不同的成长节奏，短期表现需要放进长期轨迹中理解。',True)
sub('研究观察了哪些人')
p('研究以2016—2020年当选ACM Fellow的学者为起点，筛选出科研经历至少20年的287位计算机科学家，样本主要分布于北美、欧洲和亚洲。科研经历以第一篇至最后一篇论文的发表时间跨度衡量。')
p('ACM Fellow是计算机领域的重要学术荣誉。这组样本有助于观察优秀学者的长期发展，但研究结果不能直接代表所有学科或所有科研人员。')

page()
h('01  起步快能否高产更久')
p('把科研生涯画成曲线，会发现优秀学者并没有走同一条路。仅在论文产出这一维度，研究就识别出三类典型路径。')
p('一类早期增长较慢，随后能较长时间维持高产出；一类产出上升后逐渐回落，是样本中人数最多的类型；还有一类增长较快、波动更明显，高产出持续时间相对较短。')
p('增长速度与持续时间，是理解科研成长的两个不同方面。增长最快的类型，未必拥有最长的高产出阶段。',True)
pic('image3.png',12.5)
cap('图1  三个观察维度下的成长轨迹（原稿图3）。左列为论文产出，中列为h-index，右列为p-index；不同小图对应不同类别。横轴为科研生涯年数，并非生理年龄，纵轴为标准化后的指标值。')
sub('慢一些也可能意味着较长的积累过程')
p('在综合表现维度，多数学者的p-index经过约10—20年积累达到所观察生涯中的峰值，随后虽可能回落，仍能在较长时期内保持较高水平。这里的时间尺度描述的是样本与特定指标，不能理解为人人都应遵循的成长期限。')

page()
h('02  论文变少影响力就会下降吗')
p('科研成长并不是一个数字的变化。研究分别观察论文产出、学术影响力，以及兼顾数量与引用的综合表现，并在三个维度上识别出3类、4类和6类模式。')
p('这些类别是三个维度各自形成的分类，不能相加理解为13种互不重叠的学者类型。同一个人在不同维度上，可能呈现不同节奏。',True)
sub('产出与影响力可能并不同步')
p('例如，一些学者的论文产出增长较快、随后回落，而其综合表现仍经历较长时间的积累。因此，某一阶段论文数量减少，不能直接等同于所有科研表现同步下降。')
p('学术影响力的形成同样存在差异。多数样本的h-index在发展阶段持续上升，随后逐渐趋稳；少数则经历增长与停滞交替的过程，积累更曲折，也更漫长。')
sub('同行认可也有不同的时间表')
p('研究还把当选ACM Fellow作为获得广泛同行认可的重要节点，比较不同成长类型的当选时间。')
p('从生产力和影响力维度看，一些早期产出增长较快、影响力增长较平稳的类型，在已观察科研生涯中相对较早当选；产出增长较慢或影响力发展较曲折的类型，获得认可往往相对较晚。但按p-index分类，不同群体的当选相对时间没有明显差异。')
p('这提示我们，已有成果、持续积累与获得认可之间，可能存在不同的时间关系。理解一个人的发展，需要同时看见这些过程。')
sub('三个指标分别在看什么')
p('论文产出关注发表的数量；h-index反映达到一定引用水平的论文规模；p-index综合考虑论文数量与被引次数。这些指标帮助研究者观察文献计量表现，但无法覆盖创新能力、人才培养和成果转化等全部科研贡献。')

page()
h('03  不同地区呈现怎样的成长节奏')
p('在这组ACM Fellow样本中，亚洲学者早期论文产出和h-index增长相对较快，但高产出阶段持续时间相对较短；欧洲和北美学者早期产出与影响力增长相对较缓，高产出阶段维持更久。')
p('而在p-index刻画的综合表现上，亚洲样本增长相对较慢，欧洲和北美样本相对较快。不同指标呈现的差异，再次说明了多维观察的必要性。')
pic('image6.png',11)
cap('图2  当选ACM Fellow的相对时间及成长类型的地域分布（原稿图6）。上部比较不同类别的当选相对时间，下部比较不同地区的类型构成。阅读重点是分布差异，不能据此判断个人未来。')
p('这些差异如何形成？科研政策、教育培养与资源配置等，都可能参与其中，但这项研究尚不足以确定它们各自的因果作用。',True)
p('因此，这些结果可以帮助我们提出关于科研环境的新问题。要进一步解释地区差异，还需要结合具体环境开展检验，也不宜用地区平均特征为个体贴标签。')

page()
h('04  为不同的成长节奏保留机会')
p('从这些发现出发，人才评价值得关注的，是一个人在较长时间里如何积累、变化和持续贡献。研究为完善分类支持和持续培养机制提供了参考。')
sub('把阶段性成果放回长期过程')
p('论文数量、影响力积累和表现持续性各有侧重。对积累较慢、发展较曲折的研究者，评价也需要能够反映其成长进展，减少单一时点信息带来的局限。')
sub('结合学科特点设置评价窗口')
p('科研成果往往需要时间积累。早期发展、中期积累与成熟阶段具有不同特征，评价周期应结合学科和任务特点设计，短期结果也需要放在更长的科研过程中理解。')
sub('让支持能够延续到后续成长')
p('人才支持可以兼顾早期突出者与后期突破者，为持续积累、表现不断改善的研究者保留后续机会。具体措施仍需结合科研环境设计，并通过评估检验效果。')
p('回到开头的问题：少年得志还是大器晚成？这项研究让我们看到，优秀学者并没有统一的成长时间表。更充分地理解成长，也意味着为不同节奏留下空间。',True)
sub('拓展阅读  研究如何比较成长曲线')
p('研究基于DBLP论文记录与Web of Science引用信息，估计作者贡献权重并计算指标；再用动态时间规整（DTW）比较曲线，结合主成分分析（PCA）与K-means聚类识别模式。DTW能够处理生涯长度不同、相似变化发生时间错位等情况。这里的“早期”与“成熟期”按指标峰值划分，不对应统一年龄。')
sub('论文来源')
a=p('Peng X, Xu H, Shi J. Are the bibliometric growth patterns of excellent scholars similar? From the analysis of ACM Fellows. Journal of Informetrics, 2024, 18(3): 101543.')
a.paragraph_format.line_spacing=1.15
for r in a.runs: r.font.size=Pt(9)

doc.core_properties.title='杰出学者的成长节奏'
doc.core_properties.subject='科研成长研究公众号内容排版'
doc.core_properties.author=''
for element in [doc._element, doc.styles.element]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT/'杰出学者的成长节奏_公众号初稿.docx')
print(OUT/'杰出学者的成长节奏_公众号初稿.docx')

